from __future__ import annotations

import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)


def write_sheet_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = 'A2'


def autosize_columns(ws, min_width: int = 12, max_width: int = 48) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length = max((len(str(cell.value or '')) for cell in column_cells), default=0)
        width = min(max(length + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def append_total_row(ws, *, label: str = 'Итого', numeric_start_col: int = 2) -> None:
    if ws.max_row < 2:
        return
    totals = []
    for col_idx in range(1, ws.max_column + 1):
        if col_idx < numeric_start_col:
            totals.append(label if col_idx == 1 else '')
            continue
        column_total = 0
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if isinstance(value, (int, float)):
                column_total += value
        totals.append(column_total)
    ws.append(totals)
    total_row = ws.max_row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.font = HEADER_FONT
        if col_idx >= numeric_start_col:
            cell.alignment = Alignment(horizontal='center', vertical='center')


def workbook_http_response(wb: Workbook, filename_prefix: str) -> HttpResponse:
    buffer = io.BytesIO()
    wb.save(buffer)
    timestamp = timezone.localtime().strftime('%Y-%m-%d_%H-%M')
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename_prefix}_{timestamp}.xlsx"'
    return response


def build_list_workbook(sheet_title: str, headers: list[str], rows) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    write_sheet_header(ws, headers)
    for row in rows:
        ws.append(row)
    autosize_columns(ws, max_width=56)
    return wb
