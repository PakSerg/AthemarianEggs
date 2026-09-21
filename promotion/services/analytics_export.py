from __future__ import annotations

import io
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path
from typing import BinaryIO, Iterable

from django.db.models import Count, Q
from django.db.models.functions import TruncDate
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..models import Raffle, Receipt, User
from .rejection_reasons import rejection_reasons_breakdown

DUPLICATE_RECEIPT_MESSAGE = 'Такой чек уже зарегистрирован в акции'
QR_DECODE_FAILED_MESSAGE = 'Не удалось распознать QR-код на изображении'
RECEIPT_DATE_INVALID_MESSAGE = 'Дата покупки товаров не соответствует периоду проведения акции'

RECEIPT_SHEET_HEADERS = [
    'Дата',
    'Загружено всего',
    'На проверке',
    'Прошли проверку',
    'Отклонено',
    'Победные',
    'Ошибка ФНС (ожидание ответа)',
    'ФНС: исчерпаны попытки',
    'Не распознан QR на фото',
    'Недостаточно данных для проверки',
    'ФНС не подтвердил чек',
    'Дата чека не подходит',
    'Товары не включены в акцию',
]

FRUITMOTIV_SEARCH_TERMS = ('фрутмотив', 'фр.нап', 'фрутм.нап', 'фрутм')

HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Данные за 01.07.2026-06.07.2026 не восстановить из БД, поэтому зафиксированы
# по последней корректной выгрузке (analytics_2026-07-06_14-00-37.xlsx).
HARDCODED_PARTICIPANTS_BY_DAY: dict[date, dict[str, int]] = {
    date(2026, 7, 1): {'registered': 5, 'active': 5, 'inactive': 0},
    date(2026, 7, 2): {'registered': 1, 'active': 1, 'inactive': 0},
    date(2026, 7, 3): {'registered': 1, 'active': 0, 'inactive': 1},
    date(2026, 7, 4): {'registered': 1, 'active': 1, 'inactive': 0},
    date(2026, 7, 5): {'registered': 0, 'active': 0, 'inactive': 0},
    date(2026, 7, 6): {'registered': 0, 'active': 0, 'inactive': 0},
}

_EMPTY_RECEIPT_STATS = {
    'total': 0, 'pending': 0, 'confirmed': 0, 'rejected': 0, 'winner': 0,
    'fns_waiting': 0, 'fns_exhausted': 0, 'qr_not_recognized': 0,
    'insufficient_data': 0, 'invalid_fns_response': 0,
    'receipt_date_invalid': 0, 'products_not_included': 0,
}

HARDCODED_RECEIPTS_BY_DAY: dict[date, dict[str, int]] = {
    date(2026, 7, 1): {
        **_EMPTY_RECEIPT_STATS, 'total': 3, 'confirmed': 1, 'rejected': 2, 'receipt_date_invalid': 2,
    },
    date(2026, 7, 2): {**_EMPTY_RECEIPT_STATS, 'total': 1, 'confirmed': 1},
    date(2026, 7, 3): {**_EMPTY_RECEIPT_STATS, 'total': 1, 'confirmed': 1},
    date(2026, 7, 4): {**_EMPTY_RECEIPT_STATS, 'total': 1, 'confirmed': 1},
    date(2026, 7, 5): dict(_EMPTY_RECEIPT_STATS),
    date(2026, 7, 6): dict(_EMPTY_RECEIPT_STATS),
}


def analytics_receipt_exclude_q() -> Q:
    """Чеки, не учитываемые в аналитике (дубликаты, неверный формат QR)."""
    return (
        Q(message__icontains=DUPLICATE_RECEIPT_MESSAGE)
        | Q(message__icontains='Неверный формат QR')
        | Q(system_message__icontains='Неверный формат QR')
    )


def _analytics_receipts_qs():
    return Receipt.objects.exclude(analytics_receipt_exclude_q())


def _date_range(date_from: date | None, date_to: date | None) -> list[date]:
    if date_to is None:
        date_to = timezone.localdate()
    if date_from is None:
        date_from = _detect_start_date(date_to)
    days = []
    current = date_from
    while current <= date_to:
        days.append(current)
        current += timedelta(days=1)
    return days


def _detect_start_date(date_to: date) -> date:
    candidates = []
    first_user = User.objects.order_by('created_at').values_list('created_at', flat=True).first()
    if first_user:
        candidates.append(timezone.localtime(first_user).date())
    first_receipt = Receipt.objects.order_by('created_at').values_list('created_at', flat=True).first()
    if first_receipt:
        candidates.append(timezone.localtime(first_receipt).date())
    return min(candidates) if candidates else date_to


def _default_date_range() -> tuple[date, date]:
    date_from, date_to = Raffle.get_default_analytics_range()
    if date_from is None:
        date_to = timezone.localdate()
        date_from = _detect_start_date(date_to)
    return date_from, date_to


def _write_sheet_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = 'A2'


def _autosize_columns(ws, min_width: int = 12, max_width: int = 48) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length = max(len(str(cell.value or '')) for cell in column_cells)
        width = min(max(length + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _append_total_row(ws, values: list) -> None:
    ws.append(values)
    row_idx = ws.max_row
    for col_idx, _ in enumerate(values, start=1):
        ws.cell(row=row_idx, column=col_idx).font = HEADER_FONT


def _participants_by_day(date_from: date | None, date_to: date | None) -> dict[date, dict[str, int]]:
    qs = (
        User.objects.annotate(day=TruncDate('created_at', tzinfo=timezone.get_current_timezone()))
        .values('day')
        .annotate(
            registered=Count('id'),
            active=Count('id', filter=Q(is_active=True)),
            inactive=Count('id', filter=Q(is_active=False)),
        )
    )
    if date_from:
        qs = qs.filter(day__gte=date_from)
    if date_to:
        qs = qs.filter(day__lte=date_to)

    return {
        row['day']: {
            'registered': row['registered'],
            'active': row['active'],
            'inactive': row['inactive'],
        }
        for row in qs
    }


def _receipt_error_filters() -> dict[str, Q]:
    """Фильтры для столбцов ошибок (по чекам, загруженным в этот день)."""
    return {
        'pending': Q(status=Receipt.Status.PENDING),
        'fns_waiting': (
            Q(status=Receipt.Status.PENDING)
            & (
                Q(system_message__icontains='Ошибка получения чека от ФНС')
                | Q(system_message__icontains='Ошибка получения данных чека на стороне ФНС')
                | Q(system_message__icontains='Ошибка на стороне ФНС при обработке чека')
            )
        ),
        'fns_exhausted': (
            Q(system_message__icontains='превышено количество попыток')
            | Q(message__icontains='Ошибка получения данных чека из ФНС')
        ),
        'qr_not_recognized': Q(message=QR_DECODE_FAILED_MESSAGE),
        'insufficient_data': (
            Q(message__icontains='Недостаточно данных для проверки чека')
        ),
        'invalid_fns_response': (
            Q(message__icontains='ФНС не подтвердил существование чека')
        ),
        'receipt_date_invalid': Q(message__icontains=RECEIPT_DATE_INVALID_MESSAGE),
        'products_not_included': (
            Q(status=Receipt.Status.REJECTED)
            & Q(message__icontains='Товары не соответствуют условиям акции')
        ),
    }


RECEIPT_REJECTED_BREAKDOWN_KEYS = (
    'fns_exhausted',
    'qr_not_recognized',
    'insufficient_data',
    'invalid_fns_response',
    'receipt_date_invalid',
    'products_not_included',
)

RECEIPT_REJECTED_BREAKDOWN_LABELS = {
    'fns_exhausted': 'ФНС: исчерпаны попытки',
    'qr_not_recognized': 'Не распознан QR на фото',
    'insufficient_data': 'Недостаточно данных для проверки',
    'invalid_fns_response': 'ФНС не подтвердил чек',
    'receipt_date_invalid': 'Дата чека не подходит',
    'products_not_included': 'Товары не включены в акцию',
}


def count_rejected_receipt_categories(receipts_qs) -> dict[str, int]:
    """Количество отклонённых чеков по категориям ошибок (как в Excel-аналитике)."""
    filters = _receipt_error_filters()
    annotate_kwargs = {
        key: Count('id', filter=filters[key])
        for key in RECEIPT_REJECTED_BREAKDOWN_KEYS
    }
    row = receipts_qs.filter(status=Receipt.Status.REJECTED).aggregate(**annotate_kwargs)
    return {key: row[key] for key in RECEIPT_REJECTED_BREAKDOWN_KEYS}


def _receipts_by_day(date_from: date | None, date_to: date | None) -> dict[date, dict[str, int]]:
    error_filters = _receipt_error_filters()
    annotate_kwargs = {
        'total': Count('id'),
        'confirmed': Count('id', filter=Q(status=Receipt.Status.CONFIRMED)),
        'rejected': Count('id', filter=Q(status=Receipt.Status.REJECTED)),
        'winner': Count('id', filter=Q(status=Receipt.Status.WINNER)),
    }
    for key, q_filter in error_filters.items():
        annotate_kwargs[key] = Count('id', filter=q_filter)

    qs = (
        _analytics_receipts_qs()
        .annotate(day=TruncDate('created_at', tzinfo=timezone.get_current_timezone()))
        .values('day')
        .annotate(**annotate_kwargs)
    )
    if date_from:
        qs = qs.filter(day__gte=date_from)
    if date_to:
        qs = qs.filter(day__lte=date_to)

    stat_keys = ('total', 'confirmed', 'rejected', 'winner', *error_filters.keys())
    return {
        row['day']: {key: row[key] for key in stat_keys}
        for row in qs
    }


def _receipt_row_from_stats(stats: dict) -> list:
    return [
        stats.get('total', 0),
        stats.get('pending', 0),
        stats.get('confirmed', 0),
        stats.get('rejected', 0),
        stats.get('winner', 0),
        stats.get('fns_waiting', 0),
        stats.get('fns_exhausted', 0),
        stats.get('qr_not_recognized', 0),
        stats.get('insufficient_data', 0),
        stats.get('invalid_fns_response', 0),
        stats.get('receipt_date_invalid', 0),
        stats.get('products_not_included', 0),
    ]


def _cities_stats() -> list[tuple[str, int]]:
    rows = (
        User.objects.values('city')
        .annotate(users_count=Count('id'))
        .order_by('-users_count', 'city')
    )
    result = []
    for row in rows:
        city = (row['city'] or '').strip() or 'Не указан'
        result.append((city, row['users_count']))
    return result


def _iter_receipt_items(items) -> Iterable[dict]:
    if not items:
        return
    if isinstance(items, list):
        for item in items:
            if isinstance(item, dict):
                yield item
        return
    if isinstance(items, dict):
        nested = items.get('items')
        if isinstance(nested, list):
            for item in nested:
                if isinstance(item, dict):
                    yield item
            return
        for value in items.values():
            if isinstance(value, dict) and value.get('name'):
                yield value


def _is_fruutmotiv_item(name: str) -> bool:
    normalized = name.lower()
    return any(term in normalized for term in FRUITMOTIV_SEARCH_TERMS)


def _products_stats() -> tuple[list[tuple[str, int, int]], int, int]:
    """
  Возвращает:
  - список (название, кол-во позиций, кол-во чеков)
  - всего позиций
  - всего чеков с товаром
    """
    position_counter: Counter[str] = Counter()
    receipt_ids_by_name: dict[str, set[int]] = defaultdict(set)
    receipts_with_product: set[int] = set()

    receipts = Receipt.objects.iterator(chunk_size=500)

    for receipt in receipts:
        receipt_has_match = False
        for item in _iter_receipt_items(receipt.items):
            name = (item.get('name') or '').strip()
            if not name or not _is_fruutmotiv_item(name):
                continue
            receipt_has_match = True
            position_counter[name] += 1
            receipt_ids_by_name[name].add(receipt.pk)

        if receipt_has_match:
            receipts_with_product.add(receipt.pk)

    rows = [
        (name, count, len(receipt_ids_by_name[name]))
        for name, count in position_counter.most_common()
    ]
    total_positions = sum(position_counter.values())
    return rows, total_positions, len(receipts_with_product)


def _stores_stats(receipts_qs=None) -> list[tuple[str, int]]:
    """Уникальные торговые точки (поле store в чеке) и число чеков."""
    qs = receipts_qs if receipts_qs is not None else _analytics_receipts_qs()
    counter: Counter[str] = Counter()
    for store in qs.values_list('store', flat=True).iterator(chunk_size=500):
        name = (store or '').strip() or 'Не указан'
        counter[name] += 1
    return counter.most_common()


def count_unique_stores(receipts_qs=None) -> int:
    return len(_stores_stats(receipts_qs))


def _build_workbook(date_from: date | None = None, date_to: date | None = None) -> Workbook:
    if date_from is None and date_to is None:
        date_from, date_to = _default_date_range()
    else:
        if date_to is None:
            date_to = timezone.localdate()
        if date_from is None:
            date_from = _detect_start_date(date_to)

    days = _date_range(date_from, date_to)
    participants_map = _participants_by_day(date_from, date_to)
    receipts_map = _receipts_by_day(date_from, date_to)

    wb = Workbook()

    # Лист 1 — участники
    ws_users = wb.active
    ws_users.title = 'Участники'
    _write_sheet_header(ws_users, [
        'Дата',
        'Зарегистрировано',
        'Активные (подтверждён email)',
        'Неактивные',
    ])
    for day in days:
        stats = HARDCODED_PARTICIPANTS_BY_DAY.get(day) or participants_map.get(day, {})
        ws_users.append([
            day.isoformat(),
            stats.get('registered', 0),
            stats.get('active', 0),
            stats.get('inactive', 0),
        ])
    users_total_registered = sum(
        (HARDCODED_PARTICIPANTS_BY_DAY.get(day) or participants_map.get(day, {})).get('registered', 0)
        for day in days
    )
    users_total_active = sum(
        (HARDCODED_PARTICIPANTS_BY_DAY.get(day) or participants_map.get(day, {})).get('active', 0)
        for day in days
    )
    users_total_inactive = sum(
        (HARDCODED_PARTICIPANTS_BY_DAY.get(day) or participants_map.get(day, {})).get('inactive', 0)
        for day in days
    )
    ws_users.append([])
    _append_total_row(ws_users, [
        'Итого за период',
        users_total_registered,
        users_total_active,
        users_total_inactive,
    ])
    _autosize_columns(ws_users)

    # Лист 2 — чеки
    ws_receipts = wb.create_sheet('Чеки')
    _write_sheet_header(ws_receipts, RECEIPT_SHEET_HEADERS)
    for day in days:
        stats = HARDCODED_RECEIPTS_BY_DAY.get(day) or receipts_map.get(day, {})
        ws_receipts.append([day.isoformat(), *_receipt_row_from_stats(stats)])

    receipts_totals = {key: 0 for key in _receipt_error_filters()}
    receipts_totals.update(total=0, confirmed=0, rejected=0, winner=0)
    for day in days:
        stats = HARDCODED_RECEIPTS_BY_DAY.get(day) or receipts_map.get(day, {})
        for key in receipts_totals:
            receipts_totals[key] += stats.get(key, 0)

    ws_receipts.append([])
    _append_total_row(ws_receipts, ['Итого за период', *_receipt_row_from_stats(receipts_totals)])
    _autosize_columns(ws_receipts)

    # Лист 3 — города
    ws_cities = wb.create_sheet('Города')
    _write_sheet_header(ws_cities, ['Город', 'Количество участников'])
    city_rows = _cities_stats()
    for city, count in city_rows:
        ws_cities.append([city, count])
    total_cities = len(city_rows)
    total_users_in_cities = sum((count for _, count in city_rows))
    ws_cities.append([])
    _append_total_row(ws_cities, ['Итого городов', total_cities])
    _append_total_row(ws_cities, ['Итого участников', total_users_in_cities])
    _autosize_columns(ws_cities)

    # Лист 4 — товары
    # product_rows, total_positions, total_receipts = _products_stats()
    # ws_products = wb.create_sheet('Товары Фрмотив')
    # _write_sheet_header(ws_products, [
    #     'Название товара в чеке',
    #     'Кол-во позиций',
    #     'Кол-во чеков',
    # ])
    # for name, positions, receipts_count in product_rows:
    #     ws_products.append([name, positions, receipts_count])
    # ws_products.append([])
    # _append_total_row(ws_products, ['Итого наименований', len(product_rows), ''])
    # _append_total_row(ws_products, ['Итого позиций', total_positions, ''])
    # _append_total_row(ws_products, ['Итого чеков с товаром', total_receipts, ''])
    # _autosize_columns(ws_products)

    # Лист 5 — торговые точки (только подтверждённые чеки)
    store_rows = _stores_stats(_analytics_receipts_qs().filter(status=Receipt.Status.CONFIRMED))
    ws_stores = wb.create_sheet('Торговые точки')
    _write_sheet_header(ws_stores, ['Торговая точка', 'Кол-во чеков'])
    for store_name, count in store_rows:
        ws_stores.append([store_name, count])
    total_stores = len(store_rows)
    total_store_receipts = sum(count for _, count in store_rows)
    ws_stores.append([])
    _append_total_row(ws_stores, ['Итого торговых точек', total_stores])
    _append_total_row(ws_stores, ['Итого чеков', total_store_receipts])
    _autosize_columns(ws_stores)

    # Лист 5а — причины отклонения (полная разбивка, без исключения дублей/битых QR —
    # в отличие от остальных листов этого отчёта, здесь нужна полная картина).
    receipts_in_period = Receipt.objects.all()
    if date_from:
        receipts_in_period = receipts_in_period.filter(created_at__date__gte=date_from)
    if date_to:
        receipts_in_period = receipts_in_period.filter(created_at__date__lte=date_to)
    rejection_rows = rejection_reasons_breakdown(receipts_in_period)

    ws_rejections = wb.create_sheet('Причины отклонения')
    _write_sheet_header(ws_rejections, ['Причина', 'Количество чеков', '% от отклонённых'])
    for reason in rejection_rows:
        ws_rejections.append([reason['label'], reason['count'], reason['percent']])
    if not rejection_rows:
        ws_rejections.append(['Нет отклонённых чеков за период', 0, 0])
    ws_rejections.append([])
    _append_total_row(ws_rejections, ['Итого отклонено', sum(r['count'] for r in rejection_rows), ''])
    _autosize_columns(ws_rejections)

    # Лист 6 — гарантированный приз
    from ..models import GuaranteedPrizeSent
    ws_garant = wb.create_sheet('Гарантированный приз')
    _write_sheet_header(ws_garant, ['Email участника', 'Имя', 'Фамилия', 'Телефон', 'Дата отправки', 'ID чека'])
    garant_qs = GuaranteedPrizeSent.objects.select_related('participant', 'receipt').order_by('sent_at')
    tz = timezone.get_current_timezone()
    for g in garant_qs:
        p = g.participant
        r = g.receipt
        ws_garant.append([
            p.email or '',
            p.first_name or '',
            p.last_name or '',
            p.phone or '',
            timezone.localtime(g.sent_at, tz).strftime('%d.%m.%Y %H:%M') if g.sent_at else '',
            str(r.public_id) if r else '',
        ])
    if not garant_qs.exists():
        ws_garant.append(['Нет данных'] + [''] * 5)
    _append_total_row(ws_garant, ['Итого', garant_qs.count()])
    _autosize_columns(ws_garant)

    # Лист 7 — UTM метки
    from ..models import UTMVisit
    ws_utm = wb.create_sheet('UTM метки')
    _write_sheet_header(ws_utm, ['utm_medium', 'Кол-во визитов'])
    tz = timezone.get_current_timezone()
    utm_rows = (
        UTMVisit.objects
        .values('utm_medium')
        .annotate(count=Count('id'))
        .order_by('-count')
        .exclude(utm_medium='')
    )
    for row in utm_rows:
        ws_utm.append([row['utm_medium'] or '(без метки)', row['count']])
    if not utm_rows:
        ws_utm.append(['Нет данных', 0])
    ws_utm.append([])
    _append_total_row(ws_utm, ['Итого визитов', UTMVisit.objects.exclude(utm_medium='').count()])
    _autosize_columns(ws_utm)

    # Лист 8 — победители
    from .winners_export import WINNERS_HEADERS, _collect_winner_rows
    ws_winners = wb.create_sheet('Победители')
    _write_sheet_header(ws_winners, WINNERS_HEADERS)
    winner_rows = _collect_winner_rows()
    for item in winner_rows:
        ws_winners.append(item['row'])
    if not winner_rows:
        ws_winners.append(['Нет победителей'] + [''] * (len(WINNERS_HEADERS) - 1))
    _autosize_columns(ws_winners)

    return wb


def generate_analytics_excel(
    output: str | Path | BinaryIO | None = None,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Path | bytes:
    """
    Сформировать Excel-файл с аналитикой.

    :param output: путь к файлу или file-like объект; если None — вернуть bytes.
    :param date_from: начало периода (включительно); по умолчанию — дата начала акции (Розыгрыш).
    :param date_to: конец периода (включительно); по умолчанию — дата окончания акции (Розыгрыш).
    """
    wb = _build_workbook(date_from=date_from, date_to=date_to)

    if output is None:
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    if isinstance(output, (str, Path)):
        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    wb.save(output)
    if hasattr(output, 'seek'):
        output.seek(0)
    return output
