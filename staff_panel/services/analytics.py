from __future__ import annotations

from collections import Counter
from datetime import date, timedelta
from decimal import Decimal

from django.conf import settings
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncDate
from django.utils import timezone
from openpyxl import Workbook

from promotion.models import GuaranteedPrizePayout, Raffle, Receipt, User
from promotion.services.rejection_reasons import rejection_reasons_breakdown

from .xlsx_utils import HEADER_FONT, append_total_row, autosize_columns, write_sheet_header


def _participants_qs():
    """Основная выборка участников для аналитики.

    Та же выборка, что и на странице «Участники» (см.
    staff_panel.services.participants.get_participants_queryset) — без учётных
    записей персонала (is_staff), иначе итоги по участникам в отчёте не
    совпадают со списком участников в панели.
    """
    return User.objects.all() if settings.DEBUG else User.objects.filter(is_staff=False)


def _detect_start_date(date_to: date) -> date:
    candidates = []
    first_user = _participants_qs().order_by('created_at').values_list('created_at', flat=True).first()
    if first_user:
        candidates.append(timezone.localtime(first_user).date())
    first_receipt = Receipt.objects.order_by('created_at').values_list('created_at', flat=True).first()
    if first_receipt:
        candidates.append(timezone.localtime(first_receipt).date())
    return min(candidates) if candidates else date_to


def _default_date_range() -> tuple[date, date]:
    date_from, date_to = Raffle.get_default_analytics_range()
    if date_from is None:
        date_to = timezone.localdate()
        date_from = _detect_start_date(date_to)
    else:
        # Акция официально стартует raffle.start_date, но часть участников
        # успела зарегистрироваться раньше (до объявления старта) — без этого
        # расширения дефолтный период молча исключал их из отчёта, и «Всего
        # участников зарегистрировано» не совпадало со страницей «Участники»
        # (там фильтра по дате нет вообще — это все участники за всё время).
        detected_start = _detect_start_date(date_to)
        if detected_start < date_from:
            date_from = detected_start
    return date_from, date_to


def _date_range(date_from: date | None, date_to: date | None) -> list[date]:
    if date_to is None:
        date_to = timezone.localdate()
    if date_from is None:
        date_from = _detect_start_date(date_to)
    days = []
    current = date_from
    while current <= date_to:
        days.append(current)
        current += timedelta(days=1)
    return days


def _participants_by_day(date_from: date | None, date_to: date | None) -> dict[date, dict[str, int]]:
    qs = (
        _participants_qs().annotate(day=TruncDate('created_at', tzinfo=timezone.get_current_timezone()))
        .values('day')
        .annotate(
            registered=Count('id'),
            active=Count('id', filter=Q(is_active=True)),
            inactive=Count('id', filter=Q(is_active=False)),
        )
    )
    if date_from:
        qs = qs.filter(day__gte=date_from)
    if date_to:
        qs = qs.filter(day__lte=date_to)

    return {
        row['day']: {
            'registered': row['registered'],
            'active': row['active'],
            'inactive': row['inactive'],
        }
        for row in qs
    }


def _receipts_by_day(date_from: date | None, date_to: date | None) -> dict[date, dict]:
    qs = (
        Receipt.objects.annotate(day=TruncDate('created_at', tzinfo=timezone.get_current_timezone()))
        .values('day')
        .annotate(
            pending=Count('id', filter=Q(status=Receipt.Status.PENDING)),
            confirmed=Count('id', filter=Q(status=Receipt.Status.CONFIRMED)),
            rejected=Count('id', filter=Q(status=Receipt.Status.REJECTED)),
            winner=Count('id', filter=Q(status=Receipt.Status.WINNER)),
            frozen=Count('id', filter=Q(status=Receipt.Status.FROZEN)),
            amount_sum=Sum('amount'),
        )
    )
    if date_from:
        qs = qs.filter(day__gte=date_from)
    if date_to:
        qs = qs.filter(day__lte=date_to)

    result = {}
    for row in qs:
        total = row['pending'] + row['confirmed'] + row['rejected'] + row['winner'] + row['frozen']
        result[row['day']] = {
            'total': total,
            'pending': row['pending'],
            'confirmed': row['confirmed'],
            'rejected': row['rejected'],
            'winner': row['winner'],
            'frozen': row['frozen'],
            'amount_sum': row['amount_sum'] or Decimal('0'),
        }
    return result


def _rejection_reasons_stats(date_from: date | None, date_to: date | None) -> list[dict]:
    qs = Receipt.objects.all()
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)
    return rejection_reasons_breakdown(qs)


def _cities_stats(date_from: date | None, date_to: date | None) -> list[dict]:
    """
    Разрез по городам: участники, чеки и суммы — за тот же период, что и остальные
    листы отчёта (иначе итоги по городам не сходятся с листами «Участники»/«Чеки»,
    т.к. в базе есть регистрации до старта акции).
    """
    users_qs = _participants_qs()
    if date_from:
        users_qs = users_qs.filter(created_at__date__gte=date_from)
    if date_to:
        users_qs = users_qs.filter(created_at__date__lte=date_to)

    participants_by_city: Counter[str] = Counter()
    for row in users_qs.values('city').annotate(users_count=Count('id')):
        city = (row['city'] or '').strip() or 'Не указан'
        participants_by_city[city] += row['users_count']

    receipts_qs = Receipt.objects.all()
    if date_from:
        receipts_qs = receipts_qs.filter(created_at__date__gte=date_from)
    if date_to:
        receipts_qs = receipts_qs.filter(created_at__date__lte=date_to)

    receipts_by_city: dict[str, dict] = {}
    for participant_city, amount, promo_items in (
        receipts_qs.values_list('participant__city', 'amount', 'promo_items').iterator(chunk_size=500)
    ):
        city = (participant_city or '').strip() or 'Не указан'
        bucket = receipts_by_city.setdefault(
            city, {'count': 0, 'amount': Decimal('0'), 'promo_amount': Decimal('0')},
        )
        bucket['count'] += 1
        bucket['amount'] += amount or Decimal('0')
        for item in (promo_items or []):
            if not isinstance(item, dict):
                continue
            quantity = Decimal(str(item.get('quantity') or 0))
            price = Decimal(str(item.get('price') or 0)) / 100
            bucket['promo_amount'] += price * quantity

    empty_receipts = {'count': 0, 'amount': Decimal('0'), 'promo_amount': Decimal('0')}
    all_cities = set(participants_by_city) | set(receipts_by_city)
    result = [
        {
            'name': city,
            'count': participants_by_city.get(city, 0),
            'receipts_count': receipts_by_city.get(city, empty_receipts)['count'],
            'receipts_amount': receipts_by_city.get(city, empty_receipts)['amount'],
            'promo_amount': receipts_by_city.get(city, empty_receipts)['promo_amount'],
        }
        for city in all_cities
    ]
    result.sort(key=lambda row: (-row['count'], row['name']))
    return result


def _stores_stats(date_from: date | None, date_to: date | None) -> list[tuple[str, str, int]]:
    """Разрез по магазинам: (магазин, ИНН, количество чеков)."""
    qs = Receipt.objects.filter(status__in=[Receipt.Status.CONFIRMED, Receipt.Status.WINNER])
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)

    counter: Counter[tuple[str, str]] = Counter()
    for store, inn in qs.values_list('store', 'inn').iterator(chunk_size=500):
        key = ((store or '').strip() or 'Не указан', (inn or '').strip())
        counter[key] += 1

    return sorted(
        ((store, inn, count) for (store, inn), count in counter.items()),
        key=lambda row: (-row[2], row[0].lower()),
    )


def _products_stats(date_from: date | None, date_to: date | None) -> dict:
    """
    Разрез по акционным товарам (Receipt.promo_items) за период: сводные метрики +
    таблица по каждому наименованию. Данные лежат в JSONField без готовой сводной
    таблицы — считаем в Python по всем чекам с непустыми promo_items.
    """
    qs = Receipt.objects.exclude(promo_items=[])
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)

    per_product: dict[str, dict] = {}
    receipts_with_promo = 0
    total_quantity = 0.0
    total_amount = Decimal('0')

    for promo_items in qs.values_list('promo_items', flat=True).iterator(chunk_size=500):
        if not promo_items:
            continue
        receipts_with_promo += 1
        for item in promo_items:
            if not isinstance(item, dict):
                continue
            name = (item.get('name') or '').strip() or 'Без названия'
            quantity = float(item.get('quantity') or 0)
            # price/sum в promo_items хранятся в копейках (см. items_with_promo_flags,
            # где для отображения в панели их так же делят на 100).
            price = Decimal(str(item.get('price') or 0)) / 100
            amount = price * Decimal(str(quantity))

            bucket = per_product.setdefault(name, {'name': name, 'quantity': 0.0, 'amount': Decimal('0'), 'receipts': 0})
            bucket['quantity'] += quantity
            bucket['amount'] += amount
            bucket['receipts'] += 1

            total_quantity += quantity
            total_amount += amount

    products = sorted(per_product.values(), key=lambda row: row['amount'], reverse=True)
    mode_product = max(per_product.values(), key=lambda row: row['receipts'])['name'] if per_product else '—'

    return {
        'products': products,
        'unique_products': len(per_product),
        'total_quantity': total_quantity,
        'total_amount': total_amount,
        'receipts_with_promo': receipts_with_promo,
        'avg_items_per_receipt': round(total_quantity / receipts_with_promo, 2) if receipts_with_promo else 0,
        'avg_amount_per_receipt': (
            (total_amount / receipts_with_promo) if receipts_with_promo else Decimal('0')
        ),
        'mode_product': mode_product,
        'top_products': products[:5],
    }


def _payouts_stats(date_from: date | None, date_to: date | None) -> dict:
    """
    Разрез по выплатам гарантированного приза: «создана» считается по
    created_at, «выплачена» — по paid_at, поэтому запись, созданная до начала
    периода и выплаченная внутри него, попадёт только во вторую метрику —
    это и есть смысл разреза (когда деньги реально ушли), а не дублирование.
    """
    created_qs = GuaranteedPrizePayout.objects.all()
    if date_from:
        created_qs = created_qs.filter(created_at__date__gte=date_from)
    if date_to:
        created_qs = created_qs.filter(created_at__date__lte=date_to)

    paid_qs = GuaranteedPrizePayout.objects.filter(status=GuaranteedPrizePayout.STATUS_PAID)
    if date_from:
        paid_qs = paid_qs.filter(paid_at__date__gte=date_from)
    if date_to:
        paid_qs = paid_qs.filter(paid_at__date__lte=date_to)

    status_rows = {
        row['status']: row
        for row in created_qs.values('status').annotate(count=Count('id'), amount=Sum('amount'))
    }
    status_breakdown = [
        {
            'value': value,
            'label': label,
            'count': status_rows.get(value, {}).get('count', 0),
            'amount': status_rows.get(value, {}).get('amount') or Decimal('0'),
        }
        for value, label in GuaranteedPrizePayout.STATUS_CHOICES
    ]

    total_count = created_qs.count()
    total_amount = created_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    paid_count = paid_qs.count()
    paid_amount = paid_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    # Оба терминальных статуса (ошибка и исчерпанный лимит авто-повторов)
    # одинаково означают «деньги не ушли и сама система больше не пробует».
    failed_count = created_qs.filter(
        status__in=GuaranteedPrizePayout.MANUAL_ONLY_STATUSES,
    ).count()
    hold_count = created_qs.filter(status=GuaranteedPrizePayout.STATUS_HOLD).count()

    daily_map: dict[date, dict] = {}
    tz = timezone.get_current_timezone()
    for row in created_qs.annotate(day=TruncDate('created_at', tzinfo=tz)).values('day').annotate(
        count=Count('id'), amount=Sum('amount'),
    ):
        daily_map[row['day']] = {'created_count': row['count'], 'created_amount': row['amount'] or Decimal('0')}
    for row in paid_qs.annotate(day=TruncDate('paid_at', tzinfo=tz)).values('day').annotate(
        count=Count('id'), amount=Sum('amount'),
    ):
        bucket = daily_map.setdefault(row['day'], {'created_count': 0, 'created_amount': Decimal('0')})
        bucket['paid_count'] = row['count']
        bucket['paid_amount'] = row['amount'] or Decimal('0')

    daily = []
    for day in _date_range(date_from, date_to):
        stats = daily_map.get(day, {})
        daily.append({
            'date': day,
            'created_count': stats.get('created_count', 0),
            'created_amount': stats.get('created_amount', Decimal('0')),
            'paid_count': stats.get('paid_count', 0),
            'paid_amount': stats.get('paid_amount', Decimal('0')),
        })

    durations = [
        (paid_at - created_at).total_seconds()
        for created_at, paid_at in paid_qs.values_list('created_at', 'paid_at')
        if created_at and paid_at
    ]
    avg_payout_hours = round(sum(durations) / len(durations) / 3600, 1) if durations else 0

    return {
        'total_count': total_count,
        'total_amount': total_amount,
        'paid_count': paid_count,
        'paid_amount': paid_amount,
        'failed_count': failed_count,
        'hold_count': hold_count,
        'status_breakdown': status_breakdown,
        'daily': daily,
        'avg_payout_hours': avg_payout_hours,
    }


def _funnel_stats(date_from: date | None, date_to: date | None) -> dict:
    """
    Когортная воронка: знаменатель и все шаги считаются по одной и той же группе
    участников — тем, кто зарегистрировался в выбранном периоде.
    """
    qs_u = _participants_qs()
    if date_from:
        qs_u = qs_u.filter(created_at__date__gte=date_from)
    if date_to:
        qs_u = qs_u.filter(created_at__date__lte=date_to)
    registered = qs_u.count()

    qs_r = Receipt.objects.filter(participant__in=qs_u)

    uploaded = qs_r.values('participant').distinct().count()
    confirmed = qs_r.filter(
        status__in=[Receipt.Status.CONFIRMED, Receipt.Status.WINNER]
    ).values('participant').distinct().count()
    winners = qs_r.filter(status=Receipt.Status.WINNER).values('participant').distinct().count()

    def pct(a, b):
        return round(a / b * 100) if b else 0

    return {
        'registered': registered,
        'uploaded': uploaded,
        'uploaded_pct': pct(uploaded, registered),
        'confirmed': confirmed,
        'confirmed_pct': pct(confirmed, registered),
        'winners': winners,
        'winners_pct': pct(winners, registered),
    }


def _engagement_stats(date_from: date | None, date_to: date | None) -> dict:
    qs = Receipt.objects.all()
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)

    per_user = list(qs.values('participant').annotate(n=Count('id')).values_list('n', flat=True))
    one = sum(1 for n in per_user if n == 1)
    two_three = sum(1 for n in per_user if 2 <= n <= 3)
    four_plus = sum(1 for n in per_user if n >= 4)

    users_qs = _participants_qs()
    if date_to:
        users_qs = users_qs.filter(created_at__date__lte=date_to)
    no_receipts = max(0, users_qs.count() - len(per_user))

    return {
        'no_receipts': no_receipts,
        'one': one,
        'two_three': two_three,
        'four_plus': four_plus,
    }


def _weekday_stats(date_from: date | None, date_to: date | None) -> list[dict]:
    qs = Receipt.objects.all()
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)

    day_names = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс']
    counts = [0] * 7
    tz = timezone.get_current_timezone()
    for dt in qs.values_list('created_at', flat=True).iterator(chunk_size=1000):
        counts[timezone.localtime(dt, tz).weekday()] += 1

    return [{'day': day_names[i], 'count': counts[i]} for i in range(7)]


def get_analytics_data(*, date_from: date | None = None, date_to: date | None = None) -> dict:
    if date_from is None and date_to is None:
        date_from, date_to = _default_date_range()
    elif date_to is None:
        date_to = timezone.localdate()
    elif date_from is None:
        date_from = _detect_start_date(date_to)

    days = _date_range(date_from, date_to)
    participants_map = _participants_by_day(date_from, date_to)
    receipts_map = _receipts_by_day(date_from, date_to)

    participants_daily = []
    participants_totals = {'registered': 0, 'active': 0, 'inactive': 0}
    for day in days:
        stats = participants_map.get(day, {})
        row = {
            'date': day,
            'registered': stats.get('registered', 0),
            'active': stats.get('active', 0),
            'inactive': stats.get('inactive', 0),
        }
        participants_daily.append(row)
        for key in participants_totals:
            participants_totals[key] += row[key]

    receipts_daily = []
    receipt_total_keys = ('total', 'pending', 'confirmed', 'rejected', 'winner', 'frozen')
    receipts_totals = {key: 0 for key in receipt_total_keys}
    receipts_totals['amount_sum'] = Decimal('0')
    for day in days:
        stats = receipts_map.get(day, {})
        row = {
            'date': day,
            **{key: stats.get(key, 0) for key in receipt_total_keys},
            'amount_sum': stats.get('amount_sum', Decimal('0')),
        }
        receipts_daily.append(row)
        for key in receipt_total_keys:
            receipts_totals[key] += row[key]
        receipts_totals['amount_sum'] += row['amount_sum']

    cities = _cities_stats(date_from, date_to)
    stores = [
        {'name': name, 'inn': inn, 'count': count}
        for name, inn, count in _stores_stats(date_from, date_to)
    ]

    rejection_reasons = _rejection_reasons_stats(date_from, date_to)
    products = _products_stats(date_from, date_to)
    funnel = _funnel_stats(date_from, date_to)
    engagement = _engagement_stats(date_from, date_to)
    payouts = _payouts_stats(date_from, date_to)
    weekday_activity = _weekday_stats(date_from, date_to)
    weekday_max = max((d['count'] for d in weekday_activity), default=1) or 1

    total = receipts_totals['total']
    approval_rate = round(
        (receipts_totals['confirmed'] + receipts_totals['winner']) / total * 100
    ) if total else 0
    conversion_rate = funnel['uploaded_pct']

    return {
        'date_from': date_from,
        'date_to': date_to,
        'participants_daily': participants_daily,
        'participants_totals': participants_totals,
        'receipts_daily': receipts_daily,
        'receipts_totals': receipts_totals,
        'cities': cities,
        'cities_total': len(cities),
        'stores': stores,
        'stores_total': len(stores),
        'rejection_reasons': rejection_reasons,
        'products': products,
        'funnel': funnel,
        'engagement': engagement,
        'payouts': payouts,
        'weekday_activity': weekday_activity,
        'weekday_max': weekday_max,
        'approval_rate': approval_rate,
        'conversion_rate': conversion_rate,
    }


def build_analytics_workbook(*, date_from: date | None = None, date_to: date | None = None) -> Workbook:
    data = get_analytics_data(date_from=date_from, date_to=date_to)
    date_from, date_to = data['date_from'], data['date_to']
    days = _date_range(date_from, date_to)

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = 'Сводная информация'
    write_sheet_header(ws_summary, ['Показатель', 'Значение'])
    receipts_totals = data['receipts_totals']
    participants_totals = data['participants_totals']
    products = data['products']
    top_reason = data['rejection_reasons'][0]['label'] if data['rejection_reasons'] else '—'
    summary_rows = [
        ('Период', f"{date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}"),
        ('Всего участников зарегистрировано', participants_totals['registered']),
        ('Всего чеков загружено', receipts_totals['total']),
        ('Сумма чеков, ₽', float(receipts_totals['amount_sum'] or 0)),
        ('% одобрения чеков', data['approval_rate']),
        ('Конверсия в чек, %', data['conversion_rate']),
        ('Уникальных городов', data['cities_total']),
        ('Уникальных магазинов', data['stores_total']),
        ('Чеков «На проверке»', receipts_totals['pending']),
        ('Чеков «Подтверждён»', receipts_totals['confirmed']),
        ('Чеков «Отклонён»', receipts_totals['rejected']),
        ('Чеков «Заморожен»', receipts_totals['frozen']),
        ('Чеков «Победный»', receipts_totals['winner']),
        ('Участников без чеков', data['engagement']['no_receipts']),
        ('Участников с одним чеком', data['engagement']['one']),
        ('Участников с 2-3 чеками', data['engagement']['two_three']),
        ('Участников с 4+ чеками', data['engagement']['four_plus']),
        ('Главная причина отклонения', top_reason),
        ('Сумма акционных товаров, ₽', float(products['total_amount'] or 0)),
        ('Самый частый акционный товар (мода)', products['mode_product']),
        ('Выплат гарантированного приза создано', data['payouts']['total_count']),
        ('Выплачено (статус «Выплачена»)', data['payouts']['paid_count']),
        ('Сумма выплаченного, ₽', float(data['payouts']['paid_amount'] or 0)),
        ('Выплат с ошибкой (финальной)', data['payouts']['failed_count']),
    ]
    for label, value in summary_rows:
        ws_summary.append([label, value])
    autosize_columns(ws_summary, max_width=64)

    ws_users = wb.create_sheet('Участники')
    write_sheet_header(ws_users, ['Дата', 'Зарегистрировано', 'Активные', 'Неактивные'])
    for row in data['participants_daily']:
        ws_users.append([row['date'].isoformat(), row['registered'], row['active'], row['inactive']])
    append_total_row(ws_users)
    autosize_columns(ws_users)

    ws_receipts = wb.create_sheet('Чеки')
    write_sheet_header(ws_receipts, [
        'Дата', 'Загружено всего', 'На проверке', 'Принятые', 'Отклонено', 'Заморожено', 'Победные', 'Сумма чеков',
    ])
    for row in data['receipts_daily']:
        ws_receipts.append([
            row['date'].isoformat(),
            row['total'],
            row['pending'],
            row['confirmed'],
            row['rejected'],
            row['frozen'],
            row['winner'],
            float(row['amount_sum'] or 0),
        ])
    append_total_row(ws_receipts)
    autosize_columns(ws_receipts, max_width=56)

    ws_cities = wb.create_sheet('Города')
    write_sheet_header(ws_cities, [
        'Город', 'Количество участников', 'Количество чеков', 'Сумма чеков, ₽', 'Сумма акционных товаров, ₽',
    ])
    for city in data['cities']:
        ws_cities.append([
            city['name'],
            city['count'],
            city['receipts_count'],
            float(city['receipts_amount'] or 0),
            float(city['promo_amount'] or 0),
        ])
    # Три отдельных итога вместо одной строки на все столбцы — иначе строка
    # «Итого городов» на деле суммирует «Количество участников» и вводит в
    # заблуждение (665 участников — это не число городов).
    totals_rows = [
        ['Итого городов', len(data['cities']), '', '', ''],
        ['Итого участников', sum(c['count'] for c in data['cities']), '', '', ''],
        [
            'Итого по чекам',
            '',
            sum(c['receipts_count'] for c in data['cities']),
            float(sum(c['receipts_amount'] for c in data['cities'])),
            float(sum(c['promo_amount'] for c in data['cities'])),
        ],
    ]
    for row in totals_rows:
        ws_cities.append(row)
        row_idx = ws_cities.max_row
        for col_idx in range(1, len(row) + 1):
            ws_cities.cell(row=row_idx, column=col_idx).font = HEADER_FONT
    autosize_columns(ws_cities)

    ws_stores = wb.create_sheet('Магазины')
    write_sheet_header(ws_stores, ['Магазин', 'ИНН', 'Количество чеков'])
    for store in data['stores']:
        ws_stores.append([store['name'], store['inn'], store['count']])
    append_total_row(ws_stores, label='Итого магазинов', numeric_start_col=3)
    autosize_columns(ws_stores, max_width=64)

    ws_products = wb.create_sheet('Товары')
    write_sheet_header(ws_products, ['Показатель', 'Значение'])
    ws_products.append(['Уникальных наименований', products['unique_products']])
    ws_products.append(['Всего продано единиц (акционные)', products['total_quantity']])
    ws_products.append(['Сумма акционных товаров, ₽', float(products['total_amount'] or 0)])
    ws_products.append(['Чеков с акционными товарами', products['receipts_with_promo']])
    ws_products.append(['Среднее число ед. на чек', products['avg_items_per_receipt']])
    ws_products.append(['Средняя сумма акционных товаров на чек, ₽', float(products['avg_amount_per_receipt'] or 0)])
    ws_products.append(['Мода (самый частый товар)', products['mode_product']])
    ws_products.append([])
    ws_products.append(['Товар', 'Количество', 'Сумма, ₽', 'Чеков'])
    products_header_row = ws_products.max_row
    for col_idx in range(1, 5):
        ws_products.cell(row=products_header_row, column=col_idx).font = HEADER_FONT
    for product in products['products']:
        ws_products.append([product['name'], product['quantity'], float(product['amount']), product['receipts']])
    if not products['products']:
        ws_products.append(['Нет акционных товаров за период', 0, 0, 0])
    autosize_columns(ws_products, max_width=64)

    ws_payouts = wb.create_sheet('Выплаты')
    write_sheet_header(ws_payouts, ['Показатель', 'Значение'])
    payouts = data['payouts']
    ws_payouts.append(['Всего выплат создано за период', payouts['total_count']])
    ws_payouts.append(['Сумма созданных выплат, ₽', float(payouts['total_amount'] or 0)])
    ws_payouts.append(['Выплачено (статус «Выплачена»)', payouts['paid_count']])
    ws_payouts.append(['Сумма выплаченного, ₽', float(payouts['paid_amount'] or 0)])
    ws_payouts.append(['С ошибкой (финальной)', payouts['failed_count']])
    ws_payouts.append(['Отложены (ждут запуска)', payouts['hold_count']])
    ws_payouts.append(['Среднее время до выплаты, ч', payouts['avg_payout_hours']])
    ws_payouts.append([])
    ws_payouts.append(['Статус', 'Количество', 'Сумма, ₽'])
    status_header_row = ws_payouts.max_row
    for col_idx in range(1, 4):
        ws_payouts.cell(row=status_header_row, column=col_idx).font = HEADER_FONT
    for row in payouts['status_breakdown']:
        ws_payouts.append([row['label'], row['count'], float(row['amount'] or 0)])
    ws_payouts.append([])
    ws_payouts.append(['Дата', 'Создано выплат', 'Сумма созданных, ₽', 'Выплачено', 'Сумма выплаченного, ₽'])
    daily_header_row = ws_payouts.max_row
    for col_idx in range(1, 6):
        ws_payouts.cell(row=daily_header_row, column=col_idx).font = HEADER_FONT
    for row in payouts['daily']:
        ws_payouts.append([
            row['date'].isoformat(),
            row['created_count'],
            float(row['created_amount'] or 0),
            row['paid_count'],
            float(row['paid_amount'] or 0),
        ])
    autosize_columns(ws_payouts, max_width=56)

    ws_rejections = wb.create_sheet('Причины отклонения')
    write_sheet_header(ws_rejections, ['Причина', 'Количество чеков', '% от отклонённых'])
    reasons = data['rejection_reasons']
    for reason in reasons:
        ws_rejections.append([reason['label'], reason['count'], reason['percent']])
    if not reasons:
        ws_rejections.append(['Нет отклонённых чеков за период', 0, 0])
    append_total_row(ws_rejections, label='Итого отклонено')
    autosize_columns(ws_rejections, max_width=64)

    return wb
