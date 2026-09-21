"""Реестр перечисления денежных средств Победителям Акции — генерация из
согласованного шаблона (static/legal/registry_template.xlsx, заполнен
реквизитами Заказчика/Исполнителя — см. исходный docs/Форма_Реестра_выплат_Excel.xlsx).

Один и тот же сгенерированный файл прикладывается как service_agreement ко
всем сделкам участников, попавших в этот реестр (см. PayoutRegistry) —
как при ручной точечной выплате (реестр на одного участника), так и при
плановой еженедельной рассылке (реестр на всех накопившихся за период).

Структура шаблона (лист «Реестр»):
- шапка (реквизиты сторон, период, номер) — строки 1-12, не трогаем, кроме
  D12 (номер/дата реестра);
- строка 15 — пример заполнения, подлежит удалению;
- строки 16-75 — 60 пронумерованных пустых строк под участников;
- строка 76 — ИТОГО (формулы SUM), строка 77 — количество получателей;
- строки 80+ — юридический текст и подписи, не трогаем.
"""

import logging
from io import BytesIO

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

_TEMPLATE_PATH = settings.BASE_DIR / 'static' / 'legal' / 'registry_template.xlsx'

_SHEET_NAME = 'Реестр'
_EXAMPLE_ROW = 15          # строка-пример, удаляется целиком
_FIRST_DATA_ROW = 15       # после удаления примера первая строка участников — тоже 15
TEMPLATE_DATA_ROWS = 60   # столько пустых пронумерованных строк в шаблоне

# Колонки таблицы участников (см. заголовок строки 14 в шаблоне)
_COL_NUM, _COL_LAST, _COL_FIRST, _COL_MIDDLE, _COL_STATUS, _COL_REASON, \
    _COL_AMOUNT, _COL_NDFL, _COL_PHONE, _COL_BANK = range(1, 11)


def generate_registry_xlsx(payouts, registry_number) -> bytes:
    """Заполнить шаблон реестра данными переданных выплат.

    Удаляет строку-пример и все неиспользованные пустые строки, пересчитывает
    диапазоны итоговых формул под фактическое количество участников. Заливка
    ячеек (жёлтая — заполняет Заказчик/Оператор, серая — Исполнитель) берётся
    из шаблона как есть и не меняется.
    """
    payouts = list(payouts)
    if not payouts:
        raise ValueError('Список выплат для реестра пуст')
    if len(payouts) > TEMPLATE_DATA_ROWS:
        raise ValueError(
            f'В шаблоне реестра всего {TEMPLATE_DATA_ROWS} строк под участников, '
            f'передано {len(payouts)} — нужно разбить на несколько реестров'
        )

    from ..models import SbpBank
    from .guaranteed_prize_payout import normalize_phone

    wb = load_workbook(_TEMPLATE_PATH)
    ws = wb[_SHEET_NAME]

    ws['D12'] = f'№{registry_number} от {timezone.localdate().strftime("%d.%m.%Y")}'

    # openpyxl не сдвигает объединённые ячейки при delete_rows() — снимаем все
    # объединения, запомнив их, и восстановим на пересчитанных позициях после
    # обеих операций удаления (иначе шапка/подписи ниже таблицы съедут и
    # потеряют объединение по ширине).
    original_merges = [
        (r.min_row, r.max_row, r.min_col, r.max_col) for r in list(ws.merged_cells.ranges)
    ]
    for r in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(r))

    deletions = []  # (start_row, count) — в координатах, актуальных на момент каждого удаления

    # Строка-пример (15) удаляется — реальные данные начинаются с той же строки 15.
    ws.delete_rows(_EXAMPLE_ROW, 1)
    deletions.append((_EXAMPLE_ROW, 1))

    banks_by_sbp_id = {b.sbp_id: (b.name_rus or b.name) for b in SbpBank.objects.all()}

    def _phone(payout, participant):
        """Телефон в формате СБП (11 цифр с «7»).

        `payout.phone_number` заполняется только в момент создания сделки, а
        плановый реестр формируется до выплат — тогда берём телефон из профиля
        и нормализуем его сами, чтобы формат в документе был одинаковым и для
        точечной, и для пакетной выплаты.
        """
        raw = payout.phone_number or getattr(participant, 'phone', '') or ''
        return normalize_phone(raw) or raw

    def _bank(payout, participant):
        """Наименование банка получателя.

        `payout.bank_sbp_id` тоже появляется только при создании сделки —
        до этого резолвим банк по БИК из профиля, иначе в плановом реестре
        колонка «Банк получателя» осталась бы пустой у всех участников.
        """
        sbp_id = payout.bank_sbp_id or SbpBank.resolve_sbp_id(getattr(participant, 'bank_bik', None))
        return banks_by_sbp_id.get(sbp_id, '')

    last_data_row = _FIRST_DATA_ROW + len(payouts) - 1
    for i, payout in enumerate(payouts):
        row = _FIRST_DATA_ROW + i
        participant = payout.participant
        ws.cell(row=row, column=_COL_NUM, value=i + 1)
        ws.cell(row=row, column=_COL_LAST, value=participant.last_name)
        ws.cell(row=row, column=_COL_FIRST, value=participant.first_name)
        ws.cell(row=row, column=_COL_MIDDLE, value=participant.middle_name or '')
        ws.cell(row=row, column=_COL_STATUS, value='Победитель')
        ws.cell(row=row, column=_COL_REASON, value='Гарантированный')
        ws.cell(row=row, column=_COL_AMOUNT, value=float(payout.amount))
        ws.cell(row=row, column=_COL_NDFL, value=0)
        ws.cell(row=row, column=_COL_PHONE, value=_phone(payout, participant))
        ws.cell(row=row, column=_COL_BANK, value=_bank(payout, participant))

    # Неиспользованные пустые пронумерованные строки шаблона — удаляем целиком,
    # чтобы в документе не оставалось десятков пустых строк.
    last_template_row = _FIRST_DATA_ROW + TEMPLATE_DATA_ROWS - 1  # 74 после удаления примера
    if last_data_row < last_template_row:
        count = last_template_row - last_data_row
        ws.delete_rows(last_data_row + 1, count)
        deletions.append((last_data_row + 1, count))

    # Строки «ИТОГО» / «Количество получателей» сдвинулись вплотную к данным —
    # формулы переписываем под фактический диапазон (openpyxl не пересчитывает
    # текст формул при удалении строк).
    totals_row = last_data_row + 1
    count_row = totals_row + 1
    ws.cell(row=totals_row, column=_COL_AMOUNT,
            value=f'=SUM(G{_FIRST_DATA_ROW}:G{last_data_row})')
    ws.cell(row=totals_row, column=_COL_NDFL,
            value=f'=SUM(H{_FIRST_DATA_ROW}:H{last_data_row})')
    ws.cell(row=count_row, column=_COL_AMOUNT,
            value=f'=COUNTIF(B{_FIRST_DATA_ROW}:B{last_data_row},"?*")')

    # Восстанавливаем объединения ячеек на пересчитанных позициях. Диапазоны
    # шапки/таблицы (строки 1-14) и удалённых блоков пропускаем — они либо не
    # сдвигались, либо были удалены вместе со строками.
    def _shift(row):
        for start, cnt in deletions:
            if start <= row < start + cnt:
                return None
            if row >= start + cnt:
                row -= cnt
        return row

    for min_row, max_row, min_col, max_col in original_merges:
        if min_row <= 14:
            new_min, new_max = min_row, max_row
        else:
            new_min, new_max = _shift(min_row), _shift(max_row)
            if new_min is None or new_max is None:
                continue
        ws.merge_cells(start_row=new_min, start_column=min_col, end_row=new_max, end_column=max_col)

    # Файл должен открываться на первых строках, а не там, где стоял курсор при
    # подготовке шаблона. В шаблоне закреплена шапка (freeze до строки 14), и
    # при закреплённых областях Excel прокручивает нижнюю панель к её
    # собственному pane/@topLeftCell — в шаблоне там осталась строка 83, из-за
    # чего документ открывался в самом низу. `sheet_view.topLeftCell` в этом
    # случае не действует, поэтому пере-задаём именно закрепление: шапка
    # остаётся зафиксированной, а нижняя панель начинается с первой строки
    # данных. Курсор дополнительно ставим в начало таблицы.
    ws.freeze_panes = f'A{_FIRST_DATA_ROW}'
    ws.sheet_view.topLeftCell = 'A1'
    for selection in ws.sheet_view.selection:
        selection.activeCell = f'A{_FIRST_DATA_ROW}'
        selection.sqref = f'A{_FIRST_DATA_ROW}'

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_registry(payouts, *, save_payouts=True):
    """Создать PayoutRegistry на переданные выплаты: сгенерировать xlsx по
    шаблону и сохранить его как файл реестра. Если `save_payouts` — сразу
    проставить `payout.registry` у каждой выплаты."""
    from ..models import PayoutRegistry

    registry = PayoutRegistry.objects.create()
    xlsx_bytes = generate_registry_xlsx(payouts, registry_number=str(registry.pk))
    registry.registry_number = str(registry.pk)
    registry.file.save(f'Реестр_№{registry.pk}.xlsx', ContentFile(xlsx_bytes), save=True)

    if save_payouts:
        for payout in payouts:
            payout.registry = registry
            payout.save(update_fields=['registry', 'updated_at'])

    logger.info('Создан реестр выплат №%s на %s участников', registry.pk, len(payouts))
    return registry


def create_single_participant_registry(payout):
    """Создать реестр на одного участника — используется при точечной ручной
    выплате (кнопка «Оплатить»/«Выплатить тестово»), где реестр из вопроса
    пользователя формируется «на лету», один участник на реестр."""
    return create_registry([payout])


def render_preview_xlsx() -> bytes:
    """Реестр с фиктивными данными — чтобы проверить вёрстку шаблона из
    панели персонала, не создавая реальную выплату/реестр в базе."""
    from types import SimpleNamespace

    fake_participant = SimpleNamespace(
        last_name='Иванов', first_name='Иван', middle_name='Иванович', phone='79000000001',
    )
    fake_payout = SimpleNamespace(
        participant=fake_participant, amount=settings.CYCLOPS_CONFIG['PRIZE_AMOUNT'],
        phone_number='79000000001', bank_sbp_id=None,
    )
    return generate_registry_xlsx([fake_payout], registry_number='PREVIEW')
